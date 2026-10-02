"""Trích xuất văn bản từ PDF / DOCX / TXT (giữ số trang nếu có)."""
from __future__ import annotations

import re

# Ký tự điều khiển (C0 + DEL) — pypdf đôi khi giải mã glyph bullet ("•") của PDF
# thành các byte này thay vì ký tự đọc được, làm nhiễu cả nội dung lẫn điểm truy hồi.
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


# Làm sạch văn bản: bỏ ký tự điều khiển, gộp khoảng trắng thừa, tối đa 1 dòng trống liên tiếp
def _clean(text: str) -> str:
    text = _CTRL.sub("", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# Chọn hàm trích xuất theo loại file
def extract(path: str, file_type: str) -> tuple[list[tuple[int | None, str]], int | None]:
    """Trả về (list[(page_no|None, text)], tổng_số_trang|None)."""
    ft = file_type.lower()
    if ft == "pdf":
        return _extract_pdf(path)
    if ft == "docx":
        return _extract_docx(path)
    return _extract_txt(path)


_PAGE_NO = re.compile(r"^(?:trang|page)\s*\d+(?:\s*/\s*\d+)?$|^\d+$", re.IGNORECASE)


# Bỏ tiêu đề / chân trang lặp lại (vd "Bài giảng Hệ điều hành", "Bộ môn CNPM – Khoa CNTT", "Trang 3"):
# dòng nằm ở vài dòng đầu / cuối trang và lặp lại trên >= 60% số trang, cùng các dòng chỉ ghi số trang.
# Không bỏ thì chúng chen vào giữa nội dung ở mỗi chỗ chuyển trang, làm nhiễu embedding và ngữ cảnh của AI.
def _strip_page_furniture(pages: list[str]) -> list[str]:
    edge = lambda lines: lines[:4] + lines[-3:]  # noqa: E731
    split = [[l.strip() for l in p.splitlines()] for p in pages]
    repeated: set[str] = set()
    if len(pages) >= 2:
        counts: dict[str, int] = {}
        for lines in split:
            for l in set(edge(lines)):
                counts[l] = counts.get(l, 0) + 1
        repeated = {l for l, n in counts.items() if l and n >= max(2, 0.6 * len(pages))}
    out = []
    for lines in split:
        top, bottom = set(lines[:4]), set(lines[-3:])
        kept = [l for i, l in enumerate(lines)
                if not _PAGE_NO.match(l) and not (l in repeated and (l in top or l in bottom))]
        out.append(_clean("\n".join(kept)))
    return out


# PDF: đọc từng trang bằng pypdf, giữ số trang để trích dẫn nguồn; trang lỗi / không có chữ thì bỏ qua
def _extract_pdf(path: str):
    from pypdf import PdfReader

    reader = PdfReader(path)
    raw: list[tuple[int, str]] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            txt = _clean(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            txt = ""
        if txt:
            raw.append((i, txt))
    cleaned = _strip_page_furniture([t for _, t in raw])
    out: list[tuple[int | None, str]] = [(i, t) for (i, _), t in zip(raw, cleaned) if t]
    return out, len(reader.pages)


# DOCX: lấy các đoạn văn và nội dung bảng (các ô nối bằng " | "); DOCX không có khái niệm trang
def _extract_docx(path: str):
    import docx

    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = _clean("\n".join(parts))
    return ([(None, text)] if text else []), None


# TXT/MD: đọc toàn bộ file, ký tự lỗi mã hóa được thay thế thay vì báo lỗi
def _extract_txt(path: str):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = _clean(f.read())
    return ([(None, text)] if text else []), None
