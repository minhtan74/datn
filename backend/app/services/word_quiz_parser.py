"""Đọc bộ đề trắc nghiệm từ file Word (.docx) -> bản nháp câu hỏi để giảng viên xem, sửa rồi duyệt.

Không dùng AI: đọc đúng từng câu theo định dạng quen thuộc của đề trắc nghiệm tiếng Việt, dùng được cho
mọi môn (Toán, Lý, Hóa, Sinh, Văn, Sử, Địa, Anh, Tin, GDCD...).

    Câu 1 (NB): Nội dung câu hỏi (có thể nhiều dòng, có ảnh / công thức / bảng số liệu)
    A. Phương án A          (hoặc "A) ...", "(A) ...", "a. ...", nhiều phương án trên cùng 1 dòng)
    B. Phương án B
    C. ...
    D. ...
    Đáp án: B               (tùy chọn nếu đã đánh dấu đáp án theo cách khác)
    Giải thích: ... Chọn B.  (tùy chọn; "Lời giải", "Hướng dẫn giải" cũng được)
    Độ khó: Dễ | Trung bình | Khó   (tùy chọn; nhận cả Nhận biết / Thông hiểu / Vận dụng, hoặc nhãn (NB) (TH) (VD))
    Chủ đề: ...              (tùy chọn)

Đáp án đúng nhận theo thứ tự ưu tiên:
    1. Dòng "Đáp án: X" / "Đáp án đúng là X" / "Chọn X" ngay trong câu
    2. Dấu * trước phương án ("*B. def")
    3. Bảng đáp án cuối đề ("1.B 2.C 3-A ..." hoặc bảng Word Câu | Đáp án)
    4. "Chọn X" ở cuối lời giải
    5. Định dạng: đúng 1 phương án được in đậm / gạch chân / tô màu / highlight

Các dạng đặc biệt:
    - Đoạn văn / bài đọc dùng chung ("Đọc đoạn trích sau và trả lời câu 1 đến 5", "Read the following passage...")
      được gắn vào đầu từng câu trong phạm vi.
    - Dòng hướng dẫn ("Mark the letter A, B, C, or D to indicate...") làm nội dung cho các câu không có đề riêng
      (vd câu phát âm / trọng âm của môn Anh).
    - Câu đúng / sai nhiều ý (THPT từ 2025: a) b) c) d) + "Đáp án: a) Đúng, b) Sai...") -> tách thành từng câu Đúng / Sai.
    - Chỉ số trên / dưới (H₂SO₄, m/s², 10⁻³) giữ đúng; phần gạch chân của đề (câu phát âm) được giữ.
    - Bảng số liệu trong câu -> hiển thị thành bảng.
Câu không đọc được (thiếu phương án, câu trả lời ngắn, không xác định được đáp án...) trả về trong `errors` kèm lý do;
các dòng nằm ngoài câu hỏi (tiêu đề đề thi, hướng dẫn...) trả về trong `ignored` để giảng viên kiểm tra.
"""
import io
import re
import zipfile

import docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.services.omml_latex import omml_to_latex

LETTERS = ("A", "B", "C", "D")
MAX_QUESTIONS = 200

# "Câu 1:", "Câu 1.", "Câu hỏi 1)", "Question 1:", "1." / "1)" ở đầu dòng
_Q_START = re.compile(
    r"^\s*(?:(?:câu\s*hỏi|câu|cau|question|q)\s*(\d+)\s*[:.)\-–]?|(\d+)\s*[.)])\s*(.*)$", re.IGNORECASE
)
# Nhãn mức độ / điểm ngay sau số câu: "Câu 1 (NB):", "Câu 2 [TH]", "Câu 3 (0,25 điểm):" (khớp tại vị trí cho trước)
_LEVEL_TAG = re.compile(
    r"[(\[]\s*(nb|th|vdc|vd|nhận\s*biết|thông\s*hiểu|vận\s*dụng\s*cao|vận\s*dụng)\s*[)\]]\s*[:.\-–]?\s*", re.IGNORECASE
)
_SCORE_TAG = re.compile(r"[(\[]\s*[\d.,]+\s*(?:điểm|đ)\s*[)\]]\s*[:.\-–]?\s*", re.IGNORECASE)
# Phương án ở đầu dòng: "A. ...", "a) ...", "(A) ...", "[B] ...", "*B. ...", "B* ..."
_OPT_START = re.compile(r"^\s*(\*?)\s*(?:\(([A-Da-d])\)|\[([A-Da-d])\]|([A-Da-d])(\*?)\s*[.):])\s*(.*)$")
# Các phương án tiếp theo trên CÙNG một dòng: "A. x    B. y" / "(A) x (B) y" (chỉ chữ hoa, có khoảng trắng phía trước)
_OPT_INLINE = re.compile(r"\s(\*?)(?:\(([A-D])\)|([A-D])(\*?)\s*[.)])\s+")
# Dòng đáp án: "Đáp án: B", "Đáp án đúng là B", "ĐA: B", "Chọn B", "→ Chọn đáp án B"
_ANSWER = re.compile(
    r"^\s*(?:(?:đáp\s*án|dap\s*an)(?:\s*đúng)?(?:\s*là)?|đa|answer|key|(?:→|=>)?\s*chọn(?:\s*(?:đáp|phương)\s*án)?)"
    r"\s*[:：\-–]?\s*\(?([A-Da-d])\)?\s*[.,;!]?\s*$",
    re.IGNORECASE,
)
# Dòng đáp án có nội dung khác chữ cái (câu đúng/sai nhiều ý, câu trả lời ngắn)
_ANSWER_ANY = re.compile(r"^\s*(?:đáp\s*án(?:\s*đúng)?|dap\s*an|answer|key)\s*[:：\-–]\s*(.+)$", re.IGNORECASE)
_EXPLAIN = re.compile(
    r"^\s*(?:giải\s*thích|lời\s*giải|hướng\s*dẫn(?:\s*giải)?|giai\s*thich|explanation)\s*[:：\-–]\s*(.*)$", re.IGNORECASE
)
_DIFFICULTY = re.compile(r"^\s*(?:độ\s*khó|mức\s*độ|do\s*kho|difficulty)\s*[:：\-–]\s*(.*)$", re.IGNORECASE)
_TOPIC = re.compile(r"^\s*(?:chủ\s*đề|chu\s*de|topic)\s*[:：\-–]\s*(.*)$", re.IGNORECASE)
# "Chọn B" ở cuối lời giải (chữ cái phải viết hoa để không nhầm "chọn a = 2")
_CHOOSE_IN_TEXT = re.compile(r"(?i:chọn)\s*(?:(?i:đáp\s*án|phương\s*án)\s*)?([A-D])\b")
# Tiêu đề bảng đáp án cuối đề: "ĐÁP ÁN", "BẢNG ĐÁP ÁN", "Đáp án:" đứng một mình
_KEY_HEADER = re.compile(r"^\s*(?:bảng\s*)?đáp\s*án(?:\s*đề)?\s*[:：]?\s*$", re.IGNORECASE)
# Cặp "số - đáp án" trong bảng đáp án: "1.B", "1-B", "1B", "1: B", "1)B"
_KEY_PAIR = re.compile(r"(\d+)\s*[.\-:)]?\s*([A-D])\b")
# Ý đúng / sai: "a) Đúng", "b-S", "c: Sai", "d. Đ"
_TF_PAIR = re.compile(r"([a-d])\s*[).\-:]?\s*(đúng|sai|đ|s|true|false|t|f)\b", re.IGNORECASE)
_TF_WORD = re.compile(r"đúng|sai|true|false|đ|s|t|f", re.IGNORECASE)
# Mở đầu đoạn văn / bài đọc dùng chung cho nhiều câu
_PASSAGE_START = re.compile(
    r"^\s*(?:đọc\b|read\b|dựa\s+vào|căn\s+cứ\s+vào|sử\s+dụng|use\s+the|look\s+at|"
    r"cho\s+(?:đoạn|văn\s*bản|bài|thông\s*tin))",
    re.IGNORECASE,
)
# Phạm vi câu áp dụng: "từ câu 1 đến câu 5", "câu 6-8", "questions from 27 to 34"
_RANGE = re.compile(
    r"câu\s*(\d+)\s*(?:đến|tới|-|–|~|và|,)\s*(?:câu\s*)?(\d+)|(?:questions?|from)\s*(\d+)\s*(?:to|-|–|and)\s*(\d+)",
    re.IGNORECASE,
)
# Dòng hướng dẫn làm bài
_INSTRUCTION = re.compile(
    r"^\s*(?:mark\s+the\s+letter|choose\s+the|select\s+the|chọn\s+(?:đáp\s*án|phương\s*án|câu|từ)|khoanh\s+tròn|"
    r"hãy\s+chọn|trả\s+lời\s+các\s+câu)",
    re.IGNORECASE,
)
_SECTION = re.compile(r"^\s*(?:phần|part|section)\s+(?:[ivx]+|\d+)\b", re.IGNORECASE)
# Câu cần giữ phần gạch chân (gạch chân là nội dung đề, không phải đánh dấu đáp án)
_UNDERLINE_TOPIC = re.compile(r"underlined|gạch\s*chân|gạch\s*dưới|pronunciation|phát\s*âm|stress|trọng\s*âm", re.IGNORECASE)

_DIFFICULTY_MAP = {
    "easy": ("dễ", "de", "nhận biết", "nhan biet", "easy", "cơ bản", "nb"),
    "medium": ("trung bình", "trung binh", "thông hiểu", "thong hieu", "medium", "vừa", "th"),
    "hard": ("khó", "kho", "vận dụng", "van dung", "vận dụng cao", "hard", "nâng cao", "vd", "vdc"),
}

# Namespace XML của Word (python-docx không khai báo sẵn v:, mc:)
_NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_V = "urn:schemas-microsoft-com:vml"
_NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"

# Định dạng ảnh trình duyệt hiển thị được -> đuôi file khi lưu
_IMAGE_TYPES = {
    "image/png": "png", "image/jpeg": "jpg", "image/jpg": "jpg", "image/gif": "gif",
    "image/bmp": "bmp", "image/webp": "webp", "image/svg+xml": "svg",
}
# Ký hiệu nhúng trong nội dung câu hỏi (frontend hiển thị thành ảnh / công thức / bảng)
IMG_TOKEN = "[[img:{}]]"
MATH_TOKEN = "[[math:{}]]"
TABLE_OPEN, TABLE_CLOSE = "[[table]]", "[[/table]]"
# Gạch chân: ký tự kết hợp U+0332 sau mỗi ký tự được gạch (trình duyệt tự vẽ gạch dưới)
UNDERLINE = "̲"
# Đánh dấu tạm các phần không hỗ trợ -> câu chứa chúng bị báo lỗi (không làm mất nội dung âm thầm)
_OLE_MARK = "[[ole]]"
_BAD_IMG_MARK = "[[badimg]]"

# Chỉ số trên / dưới -> ký tự Unicode (H₂SO₄, m/s², 10⁻³, Fe³⁺)
_SUP = str.maketrans("0123456789+-−=()ni", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁻⁼⁽⁾ⁿⁱ")
_SUB = str.maketrans("0123456789+-−=()aeoxhklmnpst", "₀₁₂₃₄₅₆₇₈₉₊₋₋₌₍₎ₐₑₒₓₕₖₗₘₙₚₛₜ")


class WordQuizError(Exception):
    """File không đọc được (không phải .docx, hỏng, trống...)."""


# ── Đọc văn bản kèm định dạng ─────────────────────────────────────────────────
# Mỗi ký tự đi kèm flags = (nổi bật, gạch chân); None = ký hiệu ảnh / công thức / bảng hoặc khoảng trắng thêm vào.

def _run_flags(run) -> tuple[bool, bool]:
    """(nổi bật: in đậm / highlight / chữ màu khác đen, gạch chân)."""
    strong = bool(run.bold or run.font.highlight_color)
    color = run.font.color
    if not strong and color is not None and color.type is not None:
        rgb = color.rgb
        strong = rgb is None or rgb not in (RGBColor(0, 0, 0), RGBColor(0x26, 0x26, 0x26))  # màu theme = có tô màu
    return strong, bool(run.underline)


class _Lines:
    """Gom ký tự của 1 đoạn văn thành các dòng (theo ngắt dòng mềm)."""

    def __init__(self):
        self.lines: list[list[tuple[str, tuple | None]]] = [[]]

    def add(self, text: str, flags) -> None:
        for ch in text:
            if ch in "\n\r\v":
                self.lines.append([])
            elif ch == "\t":
                self.lines[-1].extend([(" ", None)] * 4)
            else:
                self.lines[-1].append((ch, flags))

    def token(self, token: str) -> None:
        # Cách ký hiệu với chữ liền trước / sau (không thêm dấu cách đầu dòng -> không bị coi là thụt lề)
        if self.lines[-1] and not self.lines[-1][-1][0].isspace():
            self.add(" ", None)
        self.lines[-1].extend((ch, None) for ch in token)
        self.add(" ", None)


def _image_token(paragraph: Paragraph, rid: str | None, save_image) -> str:
    """Ảnh nhúng trong file (theo relationship id) -> lưu ra file -> ký hiệu [[img:url]]."""
    part = paragraph.part.related_parts.get(rid) if rid else None
    if part is None:
        return ""
    ext = _IMAGE_TYPES.get(getattr(part, "content_type", ""))
    if not ext:  # EMF / WMF / TIFF...: trình duyệt không hiển thị được
        return _BAD_IMG_MARK
    if save_image is None:
        return IMG_TOKEN.format(f"(ảnh {ext})")
    return IMG_TOKEN.format(save_image(part.blob, ext))


def _walk_media(el, paragraph: Paragraph, out: _Lines, save_image) -> None:
    """Ảnh trong <w:drawing> (DrawingML, a:blip) hoặc <w:pict> (VML cũ, v:imagedata)."""
    for blip in el.iter(f"{{{_NS_A}}}blip"):
        out.token(_image_token(paragraph, blip.get(f"{{{_NS_R}}}embed"), save_image))
    for img in el.iter(f"{{{_NS_V}}}imagedata"):
        out.token(_image_token(paragraph, img.get(f"{{{_NS_R}}}id"), save_image))


def _walk_run(r, paragraph: Paragraph, out: _Lines, save_image) -> None:
    """1 run <w:r>: chữ (kèm chỉ số trên / dưới), tab, ngắt dòng, ảnh, đối tượng nhúng (<w:object> = MathType...)."""
    run = Run(r, paragraph)
    flags = _run_flags(run)
    table = _SUP if run.font.superscript else _SUB if run.font.subscript else None
    for child in r.iterchildren():
        tag = child.tag
        if tag == qn("w:t"):
            text = child.text or ""
            out.add(text.translate(table) if table else text, flags)
        elif tag == qn("w:tab"):
            out.add("\t", flags)
        elif tag in (qn("w:br"), qn("w:cr")):
            out.add("\n", flags)
        elif tag == qn("w:object"):
            out.token(_OLE_MARK)
        elif tag == f"{{{_NS_MC}}}AlternateContent":
            # Word mới: <mc:Choice> (drawing) + <mc:Fallback> (VML) cùng 1 ảnh -> chỉ lấy Choice
            choice = child.find(f"{{{_NS_MC}}}Choice")
            _walk_media(choice if choice is not None else child, paragraph, out, save_image)
        elif tag in (qn("w:drawing"), qn("w:pict")):
            _walk_media(child, paragraph, out, save_image)


def _walk(el, paragraph: Paragraph, out: _Lines, save_image) -> None:
    """Duyệt con của đoạn văn theo đúng thứ tự: run, công thức, link, vùng sửa đổi..."""
    for child in el.iterchildren():
        tag = child.tag
        if tag == qn("w:r"):
            _walk_run(child, paragraph, out, save_image)
        elif tag == f"{{{_NS_M}}}oMath":
            out.token(MATH_TOKEN.format(omml_to_latex(child)))
        elif tag == f"{{{_NS_M}}}oMathPara":
            for omath in child.iter(f"{{{_NS_M}}}oMath"):
                out.token(MATH_TOKEN.format(omml_to_latex(omath)))
        elif tag in (qn("w:hyperlink"), qn("w:ins"), qn("w:smartTag"), qn("w:sdt"), qn("w:sdtContent"),
                     qn("w:fldSimple"), qn("w:customXml")):
            _walk(child, paragraph, out, save_image)
        # w:del (chữ đã xóa khi bật Track Changes), w:pPr, bookmark... -> bỏ qua


def _paragraph_lines(paragraph: Paragraph, save_image=None) -> list[list[tuple]]:
    """Các dòng của 1 đoạn văn; đoạn thụt lề trái (vd code) được thêm dấu cách tương ứng (~9pt / dấu cách)."""
    out = _Lines()
    left = paragraph.paragraph_format.left_indent
    indent = int(left.pt // 9) if left is not None and left.pt > 0 else 0
    _walk(paragraph._p, paragraph, out, save_image)
    return [[(" ", None)] * indent + line for line in out.lines]


def _line_text(line) -> str:
    return "".join(ch for ch, _ in line)


def _render(chars) -> str:
    """Văn bản để lưu: chữ gạch chân được thêm ký tự kết hợp U+0332 (giữ phần gạch chân của đề)."""
    return "".join(ch + UNDERLINE if f and f[1] and not ch.isspace() else ch for ch, f in chars)


def _plain_line(text: str) -> list[tuple]:
    return [(ch, None) for ch in text]


# ── Bảng trong file Word ──────────────────────────────────────────────────────

def _table_key(rows: list[list[str]]) -> dict[int, str]:
    """Bảng đáp án kẻ bảng: hàng số câu + hàng đáp án (ngang), hoặc cột số câu + cột đáp án (dọc)."""
    across, down = {}, {}
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            if not cell.isdigit():
                continue
            below = rows[r + 1][c].upper() if r + 1 < len(rows) and c < len(rows[r + 1]) else ""
            right = row[c + 1].upper() if c + 1 < len(row) else ""
            if below in LETTERS:
                down[int(cell)] = below
            if right in LETTERS:
                across[int(cell)] = right
    return down if len(down) >= len(across) else across


def _table_lines(tbl, document: DocxDocument, save_image) -> list[list[tuple]]:
    """Bảng: bảng dàn trang (chứa câu / phương án) -> đọc từng ô như đoạn văn; bảng đáp án -> dòng "1.B 2.C";
    bảng số liệu -> 1 dòng ký hiệu [[table]]...[[/table]] để hiển thị thành bảng."""
    rows_el = [list(tr.iterchildren(qn("w:tc"))) for tr in tbl.iterchildren(qn("w:tr"))]

    def cell_text(tc) -> str:
        parts = []
        for p in tc.iterchildren(qn("w:p")):
            for line in _paragraph_lines(Paragraph(p, document), save_image):
                parts.append(_line_text(line).strip())
        return " ".join(x for x in parts if x).replace("|", "/")

    rows = [[cell_text(tc) for tc in row] for row in rows_el]
    cells = [c for row in rows for c in row if c]
    if any(_Q_START.match(c) or _OPT_START.match(c) or _ANSWER.match(c) for c in cells):
        lines = []
        for tc in tbl.iter(qn("w:tc")):  # ô gộp chỉ xuất hiện 1 lần trong XML
            for p in tc.iterchildren(qn("w:p")):
                lines.extend(_paragraph_lines(Paragraph(p, document), save_image))
        return lines
    words = {c.lower() for c in cells}
    key = _table_key(rows)
    if key and words & {"câu", "câu hỏi", "đáp án", "đa", "question", "answer", "key"}:
        return [_plain_line("ĐÁP ÁN"), _plain_line(" ".join(f"{n}.{L}" for n, L in sorted(key.items())))]
    if not cells:
        return []
    body = " || ".join(" | ".join(row) for row in rows).replace("]]", "] ]")
    return [_plain_line(f"{TABLE_OPEN}{body}{TABLE_CLOSE}")]


def _document_lines(document: DocxDocument, save_image=None) -> list[list[tuple]]:
    """Mọi dòng theo đúng thứ tự trong file (đoạn văn + bảng), bỏ dòng trống."""
    out = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            lines = _paragraph_lines(Paragraph(child, document), save_image)
        elif child.tag == qn("w:tbl"):
            lines = _table_lines(child, document, save_image)
        else:
            continue
        out.extend(line for line in lines if _line_text(line).strip())
    return out


# ── Phân tích từng dòng ───────────────────────────────────────────────────────

def _ratio(chars, idx: int) -> bool:
    """Phần lớn (>= 60%) ký tự thật của phương án có định dạng thứ idx (0 = nổi bật, 1 = gạch chân)."""
    flags = [f[idx] for ch, f in chars if f is not None and not ch.isspace()]
    return bool(flags) and sum(flags) / len(flags) >= 0.6


def _split_options(line) -> list[dict] | None:
    """Dòng bắt đầu bằng phương án -> tách thành 1 hoặc nhiều phương án; không phải thì None."""
    text = _line_text(line)
    m = _OPT_START.match(text)
    if not m:
        return None
    letter = m.group(2) or m.group(3) or m.group(4)
    # (vị trí bắt đầu phương án, vị trí bắt đầu nội dung, chữ cái như trong file, có dấu *)
    starts = [(m.start(), m.start(6), letter, bool(m.group(1) or m.group(5)))]
    for im in _OPT_INLINE.finditer(text, m.start(6)):
        nxt = im.group(2) or im.group(3)
        # Chỉ nhận phương án tiếp theo đúng thứ tự (A -> B -> C -> D), tránh cắt nhầm chữ "D." trong nội dung
        if LETTERS.index(nxt) == LETTERS.index(starts[-1][2].upper()) + 1:
            starts.append((im.start() + 1, im.end(), nxt, bool(im.group(1) or im.group(4))))
    options = []
    for i, (_, body_start, written, star) in enumerate(starts):
        body_end = starts[i + 1][0] if i + 1 < len(starts) else len(text)
        chars = line[body_start:body_end]
        body = _render(chars).strip()
        options.append({
            "letter": written.upper(), "upper": written.isupper(), "text": body, "star": star,
            "strong": _ratio(chars, 0), "under": _ratio(chars, 1), "raw": f"{written}) {body}",
        })
    return options


def _difficulty(value: str) -> str | None:
    v = value.strip().lower()
    for key in ("medium", "easy", "hard"):
        if any(v == w or v.startswith(w + " ") or (len(w) > 3 and v.startswith(w)) for w in _DIFFICULTY_MAP[key]):
            return key
    return None


def _tidy(text: str, keep_indent: bool = False) -> str:
    """Bỏ khoảng trắng thừa sát ký hiệu ảnh / công thức và cuối dòng; giữ thụt lề đầu dòng khi cần (code)."""
    text = re.sub(r"[ ]{2,}(?=\[\[)", " ", text)
    text = re.sub(r"(?<=\]\])[ ]{2,}", " ", text)
    lines = [line.rstrip() if keep_indent else line.strip() for line in text.split("\n")]
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()
    if lines and keep_indent:
        lines[0] = lines[0].lstrip()
    return "\n".join(lines)


def _is_key_header(text: str) -> bool:
    """Tiêu đề bảng đáp án: "BẢNG ĐÁP ÁN" hoặc "ĐÁP ÁN" viết hoa đứng một mình (tránh nhầm dòng "Đáp án:" trong câu)."""
    return bool(_KEY_HEADER.match(text)) and ("bảng" in text.lower() or text.strip() == text.strip().upper())


def _is_key_line(text: str, min_pairs: int = 2) -> bool:
    """Dòng chỉ gồm các cặp "số-đáp án", vd "1.B 2.C 3-A" (mặc định >= 2 cặp để không nhầm câu có "1 A ... 2 B";
    ngay sau tiêu đề ĐÁP ÁN thì 1 cặp cũng nhận)."""
    upper = text.upper()
    if len(_KEY_PAIR.findall(upper)) < min_pairs:
        return False
    return not re.sub(r"[\s,;.|/\-]", "", _KEY_PAIR.sub("", upper))


def _is_heading(text: str) -> bool:
    """Tiêu đề phần / tiêu đề đề thi: "PHẦN I. ĐỌC HIỂU", "Part 2", dòng viết hoa ngắn."""
    if _SECTION.match(text):
        return True
    letters = [c for c in text if c.isalpha()]
    return len(letters) >= 4 and text == text.upper() and len(text) <= 100 and not _OPT_START.match(text)


def _answer_key(lines: list[str]) -> dict[int, str]:
    """Bảng đáp án cuối đề: dòng có từ 2 cặp "số-đáp án" trở lên, hoặc các dòng sau tiêu đề "ĐÁP ÁN"."""
    key: dict[int, str] = {}
    in_key = False
    for text in lines:
        if _is_key_header(text):
            in_key = True
            continue
        if _is_key_line(text, min_pairs=1 if in_key else 2):
            for num, letter in _KEY_PAIR.findall(text.upper()):
                key[int(num)] = letter
    return key


def _range(text: str) -> tuple[int | None, int | None]:
    m = _RANGE.search(text)
    if not m:
        # Chỉ nhắc 1 câu: "trả lời câu 3", "answer question 12" -> phạm vi đúng 1 câu
        single = re.search(r"(?:câu|questions?)\s*(\d+)\b", text, re.IGNORECASE)
        return (int(single.group(1)),) * 2 if single else (None, None)
    lo, hi = (int(x) for x in (m.group(1) or m.group(3), m.group(2) or m.group(4)))
    return min(lo, hi), max(lo, hi)


def _true_false(raw: str, count: int) -> list[bool] | None:
    """Đáp án câu đúng / sai nhiều ý: "a) Đúng, b) Sai, ..." hoặc "Đ S Đ S" -> [True, False, ...]."""
    pairs = _TF_PAIR.findall(raw)
    if len(pairs) == count:
        by_letter = {p[0].lower(): p[1].lower() in ("đúng", "đ", "true", "t") for p in pairs}
        if sorted(by_letter) == [c for c in "abcd"][:count]:
            return [by_letter[c] for c in "abcd"[:count]]
    words = _TF_WORD.findall(raw)
    if len(words) == count and not re.sub(r"(?i)đúng|sai|true|false|đ|s|t|f|[\s,;/\-]", "", raw):
        return [w.lower() in ("đúng", "đ", "true", "t") for w in words]
    return None


# ── Hàm chính ───────────────────────────────────────────────────────────────

def parse_docx(data: bytes, save_image=None) -> dict:
    """Đọc file .docx -> {"questions", "errors", "found", "truncated", "ignored"}.

    save_image(blob, ext) -> url: lưu ảnh trong câu hỏi ra file và trả URL để chèn ký hiệu [[img:url]]
    (None: không lưu, chỉ đánh dấu vị trí ảnh — dùng khi test).
    """
    if not data:
        raise WordQuizError("File trống.")
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise WordQuizError("File không phải định dạng .docx (file .doc cũ: mở bằng Word rồi Lưu thành .docx).")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # file zip nhưng không phải Word
        raise WordQuizError("Không đọc được file Word, file có thể bị hỏng.") from exc

    lines = _document_lines(document, save_image)
    texts = [_line_text(line) for line in lines]
    key = _answer_key(texts)

    blocks: list[dict] = []
    ignored: list[str] = []
    cur: dict | None = None
    collecting: dict | None = None  # đoạn văn dùng chung đang gom (chưa gặp câu đầu tiên của nó)
    active: dict | None = None      # đoạn văn / hướng dẫn đang áp dụng cho các câu tiếp theo
    in_key_section = False

    def complete() -> bool:
        return cur is not None and (len(cur["options"]) >= 2 or bool(cur["answer"]) or bool(cur["answer_raw"]))

    for line, text in zip(lines, texts):
        stripped = text.strip()
        # Hết phần câu hỏi khi gặp bảng đáp án cuối đề
        if _is_key_header(stripped) or (_is_key_line(stripped) and not _OPT_START.match(stripped)):
            in_key_section = True
            continue
        if in_key_section:
            continue

        # Các dòng thông tin kèm theo của câu hiện tại
        if cur is not None and collecting is None:
            if m := _ANSWER.match(stripped):
                cur["answer"] = m.group(1).upper()
                cur["last"] = "answer"
                continue
            if m := _ANSWER_ANY.match(stripped):
                cur["answer_raw"] = m.group(1).strip()
                cur["last"] = "answer"
                continue
            if m := _EXPLAIN.match(stripped):
                cur["explanation"] = m.group(1).strip()
                cur["last"] = "explanation"
                continue
            if m := _DIFFICULTY.match(stripped):
                cur["difficulty"] = _difficulty(m.group(1)) or cur["difficulty"]
                cur["last"] = "meta"
                continue
            if m := _TOPIC.match(stripped):
                cur["topic"] = m.group(1).strip()[:100]
                cur["last"] = "meta"
                continue

        rendered = _render(line)
        options = _split_options(line)
        q_match = _Q_START.match(text) if not options else None
        if q_match and q_match.group(2):
            num = int(q_match.group(2))
            # Dòng chỉ có số ("1. ...") chỉ là câu mới khi câu trước đã có phương án (tránh cắt nhầm các ý
            # "1) ... 2) ..." trong đề) và không nằm giữa đoạn văn đang gom (trừ khi đúng số câu đầu của đoạn)
            if cur is not None and not cur["options"] and collecting is None:
                q_match = None
            elif collecting is not None and collecting["lo"] != num:
                q_match = None
        if q_match:
            num = int(q_match.group(1) or q_match.group(2))
            start = q_match.start(3)
            difficulty = None
            for _ in range(2):  # nhãn mức độ và điểm có thể đi cùng nhau, thứ tự bất kỳ
                if m := _LEVEL_TAG.match(text, start):
                    difficulty = _difficulty(m.group(1))
                    start = m.end()
                if m := _SCORE_TAG.match(text, start):
                    start = m.end()
            rest = _render(line[start:]).strip()
            if collecting is not None:
                active, collecting = collecting, None
            # Đoạn văn có phạm vi: chỉ áp dụng cho các câu trong phạm vi
            if active and active["hi"] is not None and not (active["lo"] <= num <= active["hi"]):
                active = None if num > active["hi"] else active
            cur = {"number": num, "content": [rest], "options": [], "answer": None, "answer_raw": None,
                   "explanation": "", "difficulty": difficulty, "topic": "", "last": "content",
                   "context": active if active and (active["hi"] is None or active["lo"] <= num) else None}
            blocks.append(cur)
            continue

        # Phương án của câu hiện tại
        if options and cur is not None and collecting is None:
            # Các ý a) b) c) d) trước phương án A–D là nội dung đề (vd "Có bao nhiêu phát biểu đúng?")
            if options[0]["upper"] and cur["options"] and not any(o["upper"] for o in cur["options"]):
                cur["content"].extend(o["raw"] for o in cur["options"])
                cur["options"] = []
            cur["options"].extend(options)
            cur["last"] = "option"
            continue

        # Câu hiện tại mới có các ý a) b) c)... chưa có đáp án mà gặp dòng thường -> các ý đó là nội dung đề
        if (cur is not None and collecting is None and cur["options"] and not cur["answer"] and not cur["answer_raw"]
                and not any(o["upper"] for o in cur["options"]) and cur["last"] == "option"
                and not _is_heading(stripped) and not _PASSAGE_START.match(stripped)):
            cur["content"].extend(o["raw"] for o in cur["options"])
            cur["content"].append(rendered.rstrip())
            cur["options"] = []
            cur["last"] = "content"
            continue

        # Ngoài câu hỏi (trước câu đầu tiên, hoặc câu hiện tại đã đủ phương án / đáp án)
        if cur is None or complete():
            if _is_heading(stripped) and not (collecting is not None and not _SECTION.match(stripped)):
                collecting = active = None
                ignored.append(stripped)
                continue
            lo, hi = _range(stripped)
            if _PASSAGE_START.match(stripped) and (re.match(r"(?i)\s*(đọc|read)\b", stripped) or lo is not None):
                collecting = {"kind": "passage", "head": stripped, "lines": [], "lo": lo, "hi": hi}
                active = None
                continue
            if collecting is not None:
                collecting["lines"].append(rendered.rstrip())
                continue
            if _INSTRUCTION.match(stripped):
                active = {"kind": "instruction", "head": rendered.strip(), "lines": [], "lo": lo, "hi": hi}
                continue
            if cur is not None and cur["last"] == "explanation":
                cur["explanation"] += "\n" + stripped
                continue
            if cur is not None and cur["last"] == "option" and stripped[:1].islower():
                cur["options"][-1]["text"] += " " + rendered.strip()  # phương án xuống dòng
                continue
            ignored.append(stripped)
            continue

        # Câu hiện tại chưa đủ phương án: nối vào phần đang đọc dở
        if cur["last"] == "option" and cur["options"]:
            cur["options"][-1]["text"] += " " + rendered.strip()
        elif cur["last"] == "explanation":
            cur["explanation"] += "\n" + stripped
        else:
            cur["content"].append(rendered.rstrip())
            cur["last"] = "content"

    questions, errors = [], []
    for b in blocks[:MAX_QUESTIONS]:
        stem = _tidy("\n".join(b["content"]), keep_indent=True)
        ctx = b["context"]
        # Đoạn văn dùng chung: gắn vào đầu mọi câu; hướng dẫn: chỉ làm nội dung cho câu không có đề riêng
        if ctx and ctx["kind"] == "passage" and ctx["lines"]:
            stem = _tidy("\n".join(ctx["lines"]), keep_indent=True) + ("\n\n" + stem if stem else "")
        elif ctx and ctx["kind"] == "instruction" and not stem:
            stem = ctx["head"]
        snippet = stem.replace("\n", " ").replace(UNDERLINE, "")[-80:]
        opts = {o["letter"]: o for o in b["options"]}
        letters = [o["letter"] for o in b["options"]]
        explanation = b["explanation"]

        def fail(reason: str) -> None:
            errors.append({"number": b["number"], "snippet": snippet, "reason": reason})

        if not stem:
            fail("Thiếu nội dung câu hỏi.")
            continue
        # Phần không hỗ trợ -> báo lỗi thay vì âm thầm làm mất công thức / ảnh
        all_text = stem + " ".join(o["text"] for o in b["options"]) + explanation
        if _OLE_MARK in all_text:
            fail("Có công thức MathType / Equation 3.0 (đối tượng nhúng) — chưa hỗ trợ. Trong Word: "
                 "MathType → Convert Equations → Office Math, hoặc gõ lại bằng Insert → Equation.")
            continue
        if _BAD_IMG_MARK in all_text:
            fail("Có ảnh định dạng EMF / WMF — trình duyệt không hiển thị được, hãy chèn lại ảnh PNG hoặc JPG.")
            continue
        if len(letters) != len(set(letters)):
            fail("Có phương án bị lặp chữ cái (vd hai dòng cùng là B).")
            continue

        # Câu đúng / sai nhiều ý (a) b) c) d) + "Đáp án: a) Đúng, b) Sai...") -> tách thành từng câu Đúng / Sai
        lower_only = b["options"] and not any(o["upper"] for o in b["options"])
        if lower_only and not b["answer"]:
            tf = _true_false(b["answer_raw"] or "", len(letters))
            if tf is None:
                fail("Câu đúng / sai nhiều ý: cần dòng \"Đáp án: a) Đúng, b) Sai, c) Đúng, d) Sai\" (hoặc \"Đ S Đ S\").")
                continue
            for o, is_true in zip(b["options"], tf):
                questions.append({
                    "number": f"{b['number']}{o['letter'].lower()}",
                    "question": f"{stem}\n{o['letter'].lower()}) {_tidy(o['text'])}\n(Ý trên Đúng hay Sai?)",
                    "options": {"A": "Đúng", "B": "Sai", "C": "", "D": ""},
                    "correct_answer": "A" if is_true else "B",
                    "answer_source": "đáp án đúng / sai",
                    "explanation": _tidy(explanation).replace(UNDERLINE, ""),
                    "difficulty": b["difficulty"],
                    "topic": b["topic"] or None,
                })
            continue

        if len(letters) < 2 or "A" not in opts or "B" not in opts:
            if not letters and b["answer_raw"]:
                fail("Câu trả lời ngắn (đáp án là số / chữ, không có phương án A–D) — hệ thống chỉ hỗ trợ câu có phương án.")
            else:
                fail("Cần ít nhất 2 phương án A và B.")
            continue
        if any(not opts[L]["text"] for L in letters):
            fail("Có phương án để trống nội dung.")
            continue

        # Gạch chân là nội dung đề (câu phát âm / trọng âm / tìm lỗi sai) -> giữ, và không coi là đánh dấu đáp án
        underline_topic = bool(_UNDERLINE_TOPIC.search(stem.replace(UNDERLINE, "")))

        # Xác định đáp án đúng theo thứ tự ưu tiên (xem docstring đầu file)
        answer, source = b["answer"], "dòng Đáp án"
        if not answer:
            starred = [L for L in letters if opts[L]["star"]]
            if len(starred) == 1:
                answer, source = starred[0], "dấu *"
        if not answer and b["number"] in key:
            answer, source = key[b["number"]], "bảng đáp án"
        if not answer:
            chosen = _CHOOSE_IN_TEXT.findall(explanation)
            if chosen:
                answer, source = chosen[-1], "\"Chọn …\" trong lời giải"
        if not answer:
            marked = [L for L in letters if opts[L]["strong"] or (opts[L]["under"] and not underline_topic)]
            if len(marked) == 1:
                answer, source = marked[0], "định dạng in đậm / gạch chân / tô màu"
        if not answer:
            fail("Không xác định được đáp án đúng (thêm dòng \"Đáp án: X\" hoặc in đậm phương án đúng).")
            continue
        if answer not in opts:
            fail(f"Đáp án là {answer} nhưng câu không có phương án {answer}.")
            continue

        def option_text(L: str) -> str:
            text = _tidy(opts[L]["text"]) if L in opts else ""
            return text if underline_topic else text.replace(UNDERLINE, "")

        questions.append({
            "number": b["number"],
            "question": stem,
            "options": {L: option_text(L) for L in LETTERS},
            "correct_answer": answer,
            "answer_source": source,
            "explanation": _tidy(explanation).replace(UNDERLINE, ""),
            "difficulty": b["difficulty"],
            "topic": b["topic"] or None,
        })
    return {
        "questions": questions, "errors": errors, "found": len(blocks), "truncated": len(blocks) > MAX_QUESTIONS,
        "ignored": ignored[:30], "ignored_count": len(ignored),
    }


def build_template() -> bytes:
    """File Word mẫu cho giảng viên tải về và điền theo."""
    d = docx.Document()
    d.add_heading("ĐỀ TRẮC NGHIỆM MẪU — StudyOnline", level=1)
    d.add_paragraph(
        "Hướng dẫn: mỗi câu bắt đầu bằng \"Câu <số>:\", các phương án A–D mỗi phương án một dòng "
        "(hoặc cùng một dòng). Đánh dấu đáp án đúng bằng dòng \"Đáp án: X\", HOẶC in đậm / tô màu "
        "phương án đúng, HOẶC thêm dấu * trước phương án đúng, HOẶC ghi bảng đáp án ở cuối đề. "
        "Các dòng Giải thích, Độ khó (Dễ / Trung bình / Khó), Chủ đề là tùy chọn. Xóa đoạn hướng dẫn này trước khi nhập."
    )
    d.add_paragraph("Câu 1 (NB): Từ khóa nào dùng để định nghĩa hàm trong Python?")
    for t in ("A. func", "B. def", "C. function", "D. lambda"):
        d.add_paragraph(t)
    d.add_paragraph("Đáp án: B")
    d.add_paragraph("Giải thích: def là từ khóa khai báo hàm trong Python.")
    d.add_paragraph("Chủ đề: Hàm")

    d.add_paragraph("Câu 2 (TH): Kiểu dữ liệu nào có thể thay đổi (mutable)?")
    for letter, text in (("A", "tuple"), ("B", "str"), ("C", "list"), ("D", "int")):
        p = d.add_paragraph()
        run = p.add_run(f"{letter}. {text}")
        run.bold = letter == "C"  # đáp án đúng in đậm

    d.add_paragraph("Câu 3: Kết quả của 2 ** 3 là bao nhiêu?")
    d.add_paragraph("A. 6        B. 8        C. 9        D. 5")

    d.add_paragraph("Đọc đoạn văn sau và trả lời các câu từ câu 4 đến câu 5:")
    d.add_paragraph("Python là ngôn ngữ lập trình bậc cao do Guido van Rossum tạo ra, phát hành lần đầu năm 1991.")
    d.add_paragraph("Câu 4: Ai tạo ra Python?")
    d.add_paragraph("A. Guido van Rossum     B. James Gosling     C. Dennis Ritchie     D. Bjarne Stroustrup")
    d.add_paragraph("Đáp án: A")
    d.add_paragraph("Câu 5: Python phát hành lần đầu năm nào?")
    d.add_paragraph("A. 1985     B. 1991     C. 2000     D. 2008")
    d.add_paragraph("Đáp án: B")

    d.add_paragraph("Câu 6: Xét tính đúng sai của các phát biểu về Python:")
    for t in ("a) Python là ngôn ngữ thông dịch.", "b) Python bắt buộc khai báo kiểu biến.",
              "c) Python dùng thụt lề để tạo khối lệnh."):
        d.add_paragraph(t)
    d.add_paragraph("Đáp án: a) Đúng, b) Sai, c) Đúng")

    d.add_heading("ĐÁP ÁN", level=2)
    d.add_paragraph("3.B")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
