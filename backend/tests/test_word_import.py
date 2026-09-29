r"""Nhập bộ đề trắc nghiệm từ file Word (.docx): bộ đọc + API.

    cd backend
    .\.venv\Scripts\python.exe -m pytest tests/test_word_import.py -q
"""
import io
import os
import sys

import docx
import pytest
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.main import app  # noqa: E402
from app.services.word_quiz_parser import WordQuizError, build_template, parse_docx  # noqa: E402

client = TestClient(app)


# Tạo file .docx từ danh sách dòng; mỗi dòng là chuỗi hoặc list các (chữ, kiểu) với kiểu:
# "b" in đậm, "u" gạch chân, "red" chữ đỏ, "hl" highlight
def make_docx(lines, table=None) -> bytes:
    d = docx.Document()
    for line in lines:
        p = d.add_paragraph()
        for text, style in ([(line, "")] if isinstance(line, str) else line):
            run = p.add_run(text)
            run.bold = "b" in style.split()
            run.underline = "u" in style.split()
            if "red" in style.split():
                run.font.color.rgb = RGBColor(0xFF, 0, 0)
            if "hl" in style.split():
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, text in enumerate(row):
                t.cell(r, c).text = text
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def answers(result):
    return {q["number"]: q["correct_answer"] for q in result["questions"]}


# ── Bộ đọc ──────────────────────────────────────────────────────────────────

def test_answer_line_explanation_difficulty_topic():
    r = parse_docx(make_docx([
        "ĐỀ KIỂM TRA 15 PHÚT",  # tiêu đề trước câu 1 bị bỏ qua
        "Câu 1: Từ khóa nào định nghĩa hàm?",
        "A. func", "B. def", "C. function", "D. lambda",
        "Đáp án: B",
        "Giải thích: def khai báo hàm.",
        "Độ khó: Nhận biết",
        "Chủ đề: Hàm",
    ]))
    q = r["questions"][0]
    assert r["errors"] == [] and q["correct_answer"] == "B" and q["answer_source"] == "dòng Đáp án"
    assert q["options"] == {"A": "func", "B": "def", "C": "function", "D": "lambda"}
    assert (q["explanation"], q["difficulty"], q["topic"]) == ("def khai báo hàm.", "easy", "Hàm")


@pytest.mark.parametrize("style", ["b", "u", "red", "hl"])
def test_answer_by_formatting(style):
    r = parse_docx(make_docx([
        [("Câu 1: ", "b"), ("Kiểu nào mutable?", "")],  # tiêu đề câu in đậm không ảnh hưởng
        "A. tuple", "B. str", [("C. list", style)], "D. int",
    ]))
    assert answers(r) == {1: "C"} and "định dạng" in r["questions"][0]["answer_source"]


def test_answer_by_star_and_inline_options_and_lowercase():
    r = parse_docx(make_docx([
        "Câu 1: 2 ** 3 = ?",
        "A. 6        *B. 8        C. 9        D. 5",
        "Câu 2: Chọn đúng",
        "a) sai", "b) đúng",
        "Đáp án: b",
    ]))
    assert answers(r) == {1: "B", 2: "B"}
    assert r["questions"][0]["options"] == {"A": "6", "B": "8", "C": "9", "D": "5"}
    assert r["questions"][1]["options"] == {"A": "sai", "B": "đúng", "C": "", "D": ""}


def test_answer_key_at_end_and_multiline_question():
    r = parse_docx(make_docx([
        "Câu 1: Cho đoạn code:",
        "x = [1, 2]",
        "Cho các mệnh đề: 1) x là list 2) x có 2 phần tử. Mệnh đề nào đúng?",  # "1)" không bị cắt thành câu mới
        "A. Chỉ 1", "B. Chỉ 2", "C. Cả hai", "D. Không cái nào",
        "Câu 2: 1 + 1 = ?", "A. 1", "B. 2",
        "ĐÁP ÁN",
        "1.C 2.B",
    ]))
    assert answers(r) == {1: "C", 2: "B"} and r["found"] == 2
    assert r["questions"][0]["question"].count("\n") == 2 and "x = [1, 2]" in r["questions"][0]["question"]


def test_numbered_questions_and_table_layout():
    r = parse_docx(make_docx(["1. Thủ đô Việt Nam?", "A. Hà Nội", "B. Huế", "Đáp án: A"],
                             table=[["2. Số nguyên tố nhỏ nhất?", ""], ["A. 1", "B. 2"], ["Đáp án: B", ""]]))
    assert answers(r) == {1: "A", 2: "B"}


def test_errors_reported_with_reason():
    r = parse_docx(make_docx([
        "Câu 1: Thiếu đáp án", "A. x", "B. y",
        "Câu 2: Chỉ 1 phương án", "A. x", "Đáp án: A",
        "Câu 3: Đáp án không tồn tại", "A. x", "B. y", "Đáp án: D",
        "Câu 4: Hai phương án cùng in đậm", [("A. x", "b")], [("B. y", "b")],
        "Câu 5: Hợp lệ", "A. x", "B. y", "Đáp án: A",
    ]))
    assert answers(r) == {5: "A"}
    reasons = {e["number"]: e["reason"] for e in r["errors"]}
    assert set(reasons) == {1, 2, 3, 4}
    assert "đáp án" in reasons[1].lower() and "2 phương án" in reasons[2] and "phương án D" in reasons[3]


def test_invalid_file():
    with pytest.raises(WordQuizError):
        parse_docx(b"")
    with pytest.raises(WordQuizError):
        parse_docx(b"day khong phai file word")


def test_template_parses_cleanly():
    r = parse_docx(build_template())
    assert r["errors"] == []
    assert answers(r) == {1: "B", 2: "C", 3: "B", 4: "A", 5: "B", "6a": "A", "6b": "B", "6c": "A"}
    # Câu 4, 5 kèm đoạn văn dùng chung; câu 1 có nhãn (NB) -> Dễ
    q = {x["number"]: x for x in r["questions"]}
    assert q[4]["question"].startswith("Python là ngôn ngữ") and q[1]["difficulty"] == "easy"


# ── Các dạng câu hỏi theo môn học ────────────────────────────────────────────

def make_rich_docx(items) -> bytes:
    """items: chuỗi = 1 đoạn; list[(chữ, kiểu)] = 1 đoạn nhiều run (kiểu: b / u / sup / sub); ("table", rows)."""
    d = docx.Document()
    for it in items:
        if isinstance(it, tuple) and it[0] == "table":
            rows = it[1]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            for i, row in enumerate(rows):
                for j, v in enumerate(row):
                    t.cell(i, j).text = v
            continue
        p = d.add_paragraph()
        for text, style in ([(it, "")] if isinstance(it, str) else it):
            run = p.add_run(text)
            run.bold = "b" in style.split()
            run.underline = "u" in style.split()
            run.font.superscript = "sup" in style.split()
            run.font.subscript = "sub" in style.split()
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_level_tag_and_choose_in_solution():  # Toán / Lý
    r = parse_docx(make_rich_docx(["Câu 1 (NB): Số nào là số nguyên tố?", "A. 4", "B. 7", "C. 9", "D. 15",
                                   "Lời giải: 7 chỉ chia hết cho 1 và chính nó. Chọn B."]))
    q = r["questions"][0]
    assert (q["question"], q["correct_answer"], q["difficulty"]) == ("Số nào là số nguyên tố?", "B", "easy")


def test_chemistry_physics_sub_superscripts():  # Hóa / Lý
    r = parse_docx(make_rich_docx([
        "Câu 1: Công thức axit sunfuric?", [("A. H", ""), ("2", "sub"), ("SO", ""), ("4", "sub")], "B. HCl",
        [("C. Fe", ""), ("3+", "sup")], [("D. 10", ""), ("-3", "sup"), (" m/s", ""), ("2", "sup")], "Đáp án: A"]))
    o = r["questions"][0]["options"]
    assert (o["A"], o["C"], o["D"]) == ("H₂SO₄", "Fe³⁺", "10⁻³ m/s²")


def test_english_pronunciation_keeps_underline_and_uses_instruction():  # Anh
    r = parse_docx(make_rich_docx([
        "Mark the letter A, B, C, or D to indicate the word whose underlined part differs from the other three in pronunciation.",
        "Question 1:", [("A. h", ""), ("ou", "u"), ("se", "")], [("B. m", ""), ("ou", "u"), ("se", "")],
        [("C. c", ""), ("ou", "u"), ("ntry", "")], [("D. ", ""), ("ou", "u"), ("t", "")], "ĐÁP ÁN", "1.C"]))
    q = r["questions"][0]
    assert q["question"].startswith("Mark the letter") and q["correct_answer"] == "C"
    assert q["options"]["A"] == "ho̲u̲se"  # phần gạch chân được giữ


def test_underlined_option_letter_still_parsed():  # đáp án đánh dấu bằng gạch chân cả dòng
    r = parse_docx(make_rich_docx(["Câu 1: Kiểu nào mutable?", "A. tuple", "B. str", [("C. list", "u")], "D. int"]))
    q = r["questions"][0]
    assert q["correct_answer"] == "C" and q["options"]["C"] == "list"  # gạch chân đánh dấu đáp án bị bỏ khi hiển thị


def test_shared_reading_passage():  # Anh / Văn
    r = parse_docx(make_rich_docx([
        "PHẦN I. ĐỌC HIỂU", "Đọc đoạn trích sau và trả lời các câu hỏi từ câu 1 đến câu 2:",
        "MÙA XUÂN NHO NHỎ", "Mọc giữa dòng sông xanh / Một bông hoa tím biếc",
        "Câu 1: Đoạn thơ viết theo thể thơ nào?", "A. Năm chữ", "B. Lục bát", "C. Tự do", "D. Bảy chữ", "Đáp án: A",
        "Câu 2: Hình ảnh nào xuất hiện?", "A. Bông hoa", "B. Con thuyền", "C. Cánh cò", "D. Vầng trăng", "Đáp án: A",
        "Read the following passage and answer the questions from 3 to 3.", "Tom lives in a small village.",
        "Question 3: Where does Tom live?", "A. City", "B. Village", "C. Town", "D. Farm", "Đáp án: B",
        "Câu 4: Câu không thuộc đoạn văn nào?", "A. x", "B. y", "Đáp án: A",
    ]))
    q = {x["number"]: x["question"] for x in r["questions"]}
    assert q[1].startswith("MÙA XUÂN NHO NHỎ\nMọc giữa dòng sông xanh") and q[1].endswith("thể thơ nào?")
    assert q[2].startswith("MÙA XUÂN NHO NHỎ") and q[3].startswith("Tom lives in a small village.")
    assert q[4] == "Câu không thuộc đoạn văn nào?"  # ngoài phạm vi 3–3 -> không gắn đoạn văn
    assert "PHẦN I. ĐỌC HIỂU" in r["ignored"]


def test_passage_for_single_question_not_leaking():  # "trả lời câu 3" chỉ áp dụng cho câu 3
    r = parse_docx(make_rich_docx([
        "Đọc đoạn thơ sau và trả lời câu 1:", "Mọc giữa dòng sông xanh",
        "Câu 1: Màu hoa?", "A. Tím", "B. Đỏ", "Đáp án: A",
        "Câu 2: 1 + 1 = ?", "A. 2", "B. 3", "Đáp án: A",
    ]))
    q = {x["number"]: x["question"] for x in r["questions"]}
    assert q[1].startswith("Mọc giữa dòng sông xanh") and q[2] == "1 + 1 = ?"


def test_statements_inside_question_then_options():  # Sinh: "có bao nhiêu phát biểu đúng"
    r = parse_docx(make_rich_docx([
        "Câu 1: Cho các phát biểu sau về ADN:", "a) ADN có cấu trúc xoắn kép.", "b) ADN chỉ có ở nhân.",
        "Có bao nhiêu phát biểu đúng?", "A. 0", "B. 1", "C. 2", "D. 3", "Đáp án: B"]))
    q = r["questions"][0]
    assert q["correct_answer"] == "B" and "a) ADN có cấu trúc xoắn kép." in q["question"]
    assert q["question"].endswith("Có bao nhiêu phát biểu đúng?")


@pytest.mark.parametrize("answer_line", ["Đáp án: a) Đúng, b) Sai, c) Đúng", "Đáp án: Đ S Đ", "Đáp án: a-Đ; b-S; c-Đ"])
def test_true_false_multi_statement_split(answer_line):  # THPT 2025 phần II
    r = parse_docx(make_rich_docx([
        "Câu 1: Xét tính đúng sai:", "a) 2 là số chẵn.", "b) 3 là số chẵn.", "c) 4 là số chẵn.", answer_line]))
    assert answers(r) == {"1a": "A", "1b": "B", "1c": "A"}
    q = r["questions"][1]
    assert q["options"]["A"] == "Đúng" and "b) 3 là số chẵn." in q["question"]


def test_true_false_without_answers_and_short_answer_reported():  # THPT 2025 phần II, III
    r = parse_docx(make_rich_docx([
        "Câu 1: Xét tính đúng sai:", "a) x", "b) y",
        "Câu 2: Tính 2 + 7.", "Đáp án: 9"]))
    reasons = {e["number"]: e["reason"] for e in r["errors"]}
    assert "đúng / sai" in reasons[1] and "trả lời ngắn" in reasons[2]


def test_data_table_in_question_and_answer_key_table():  # Địa + bảng đáp án kẻ bảng
    r = parse_docx(make_rich_docx([
        "Câu 1: Cho bảng số liệu:", ("table", [["Năm", "2010", "2020"], ["Dân số", "87,1", "97,6"]]),
        "Nhận xét nào đúng?", "A. Giảm", "B. Tăng",
        "Câu 2: 1 + 1 = ?", "A. 2", "B. 3",
        "BẢNG ĐÁP ÁN", ("table", [["Câu", "1", "2"], ["Đáp án", "B", "A"]]),
    ]))
    assert answers(r) == {1: "B", 2: "A"}
    assert "[[table]]Năm | 2010 | 2020 || Dân số | 87,1 | 97,6[[/table]]" in r["questions"][0]["question"]


def test_code_indentation_and_paren_options():  # Tin + phương án (A)..(D)
    r = parse_docx(make_rich_docx([
        "Câu 1: Đoạn code sau in ra gì?", "for i in range(2):", "    print(i)",
        "(A) 0 1", "(B) 1 2", "(C) 0 1 2", "(D) Lỗi", "Đáp án đúng là A"]))
    q = r["questions"][0]
    assert q["question"] == "Đoạn code sau in ra gì?\nfor i in range(2):\n    print(i)" and q["correct_answer"] == "A"


def test_instruction_between_questions_not_glued_to_option():  # Sử / GDCD
    r = parse_docx(make_rich_docx([
        "Câu 1: Năm 1945?", "A. 1930", "B. 1945", "Đáp án: B", "Chọn phương án đúng nhất cho các câu sau:",
        "Câu 2: Năm 1954?", "A. 1945", "B. 1954", "Đáp án: B"]))
    q = {x["number"]: x for x in r["questions"]}
    assert q[1]["options"]["B"] == "1945" and q[2]["question"] == "Năm 1954?"


# ── API ─────────────────────────────────────────────────────────────────────

def _h(email):
    tok = client.post("/api/auth/login", json={"email": email, "password": "123456"}).json()["token"]
    return {"Authorization": f"Bearer {tok}"}


def test_import_word_api_flow():
    T, S = _h("an.nguyen@studyonline.vn"), _h("em.hoang@gmail.com")
    ctype = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    # File mẫu tải được và là file Word
    tpl = client.get("/api/quizzes/import-word/template", headers=T)
    assert tpl.status_code == 200 and tpl.content[:2] == b"PK"
    # Học viên không được dùng; sai đuôi / không có câu hỏi bị từ chối
    assert client.post("/api/quizzes/import-word", headers=S,
                       files={"file": ("de.docx", tpl.content, ctype)}).status_code == 403
    assert client.post("/api/quizzes/import-word", headers=T,
                       files={"file": ("de.doc", b"x", "application/msword")}).status_code == 400
    empty = make_docx(["Chỉ là văn bản, không có câu hỏi."])
    assert client.post("/api/quizzes/import-word", headers=T,
                       files={"file": ("de.docx", empty, ctype)}).status_code == 422

    # Đọc file -> bản nháp (chưa lưu) -> lưu thành quiz mới nguồn "word" (câu 2 phương án vẫn hợp lệ)
    data = make_docx(["Câu 1: Python là?", "A. Ngôn ngữ", "B. Con rắn", "Đáp án: A"])
    got = client.post("/api/quizzes/import-word", headers=T, files={"file": ("de.docx", data, ctype)}).json()["data"]
    assert len(got["questions"]) == 1 and got["errors"] == []
    cid = int(client.post("/api/courses", headers=T, json={"title": "PYTEST Word", "description": "", "thumbnail": ""}).json()["id"])
    try:
        saved = client.post("/api/ai/quiz/approve", headers=T, json={
            "course_id": cid, "title": "Đề từ Word", "questions": got["questions"], "source": "word"}).json()
        assert saved["success"] and saved["data"]["question_count"] == 1
        quiz = client.get(f"/api/quizzes?id={saved['data']['quiz_id']}").json()["data"]
        assert quiz["source"] == "word" and quiz["question_count"] == 1
        # Nguồn AI vẫn bắt buộc đủ 4 phương án; nguồn lạ bị từ chối
        assert client.post("/api/ai/quiz/approve", headers=T, json={
            "course_id": cid, "title": "x", "questions": got["questions"]}).status_code == 400
        assert client.post("/api/ai/quiz/approve", headers=T, json={
            "course_id": cid, "title": "x", "questions": got["questions"], "source": "pdf"}).status_code == 400
    finally:
        client.request("DELETE", f"/api/courses?id={cid}", headers=T)


# ── Ảnh + công thức (Equation của Word) ──────────────────────────────────────

from docx.oxml import parse_xml  # noqa: E402

from app.services.omml_latex import omml_to_latex  # noqa: E402

_MNS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


def omath(inner: str):
    return parse_xml(f"<m:oMath {_MNS}>{inner}</m:oMath>")


def r(t: str) -> str:
    return f"<m:r><m:t>{t}</m:t></m:r>"


# Ảnh PNG 4x4 hợp lệ để chèn vào file Word (tạo bằng Python thuần, không cần Pillow)
def _png() -> bytes:
    import struct
    import zlib

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    raw = b"".join(b"\x00" + b"\x25\x63\xeb" * 4 for _ in range(4))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 4, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


PNG = _png()


@pytest.mark.parametrize("inner, expected", [
    (f"<m:f><m:num>{r('1')}</m:num><m:den>{r('2')}</m:den></m:f>", r"\frac{1}{2}"),
    (f"<m:sSup><m:e>{r('x')}</m:e><m:sup>{r('2')}</m:sup></m:sSup>{r('+1')}", r"{x}^{2}+1"),
    (f"<m:sSub><m:e>{r('a')}</m:e><m:sub>{r('n')}</m:sub></m:sSub>", r"{a}_{n}"),
    (f"<m:rad><m:radPr><m:degHide m:val='1'/></m:radPr><m:deg/><m:e>{r('x')}</m:e></m:rad>", r"\sqrt{x}"),
    (f"<m:rad><m:deg>{r('3')}</m:deg><m:e>{r('8')}</m:e></m:rad>", r"\sqrt[3]{8}"),
    (f"<m:nary><m:naryPr><m:chr m:val='∑'/></m:naryPr><m:sub>{r('i=1')}</m:sub><m:sup>{r('n')}</m:sup>"
     f"<m:e>{r('i')}</m:e></m:nary>", r"\sum_{i=1}^{n} {i}"),
    (f"<m:nary><m:sub>{r('0')}</m:sub><m:sup>{r('1')}</m:sup><m:e>{r('x dx')}</m:e></m:nary>", r"\int_{0}^{1} {x dx}"),
    (f"<m:d><m:dPr><m:begChr m:val='{{'/><m:endChr m:val=''/></m:dPr><m:e><m:eqArr><m:e>{r('x+y=2')}</m:e>"
     f"<m:e>{r('x-y=0')}</m:e></m:eqArr></m:e></m:d>", r"\left\{ \begin{array}{l} x+y=2 \\ x-y=0 \end{array} \right."),
    (f"<m:func><m:fName>{r('sin')}</m:fName><m:e>{r('x')}</m:e></m:func>", r"\sin {x}"),
    (f"<m:limLow><m:e>{r('lim')}</m:e><m:lim>{r('x→0')}</m:lim></m:limLow>", r"\lim_{x\to 0}"),
    (f"<m:acc><m:accPr><m:chr m:val='&#x20D7;'/></m:accPr><m:e>{r('AB')}</m:e></m:acc>", r"\vec{AB}"),
    (f"<m:d><m:e><m:m><m:mr><m:e>{r('1')}</m:e><m:e>{r('2')}</m:e></m:mr><m:mr><m:e>{r('3')}</m:e>"
     f"<m:e>{r('4')}</m:e></m:mr></m:m></m:e></m:d>", r"\left( \begin{matrix} 1 & 2 \\ 3 & 4 \end{matrix} \right)"),
    (r("a≤b, π≠3, x∈ℝ"), r"a\le b, \pi \ne 3, x\in \mathbb{R}"),
])
def test_omml_to_latex(inner, expected):
    assert omml_to_latex(omath(inner)) == expected


def test_image_and_equation_in_question_and_options():
    d = docx.Document()
    d.add_paragraph("Câu 1: Giá trị của biểu thức")
    p = d.paragraphs[-1]
    p._p.append(omath(f"<m:f><m:num>{r('1')}</m:num><m:den>{r('2')}</m:den></m:f>"))
    p.add_run(" là?")
    for letter, val in (("A", "0.5"), ("B", "2")):
        op = d.add_paragraph(f"{letter}. ")
        op._p.append(omath(r(val)))
    d.add_paragraph("Đáp án: A")
    d.add_paragraph("Câu 2: Hình nào là hình tròn?")
    d.add_paragraph().add_run().add_picture(io.BytesIO(PNG))
    pa = d.add_paragraph("A. ")
    pa.add_run().add_picture(io.BytesIO(PNG))
    d.add_paragraph("B. Hình vuông")
    d.add_paragraph("Đáp án: A")
    buf = io.BytesIO()
    d.save(buf)

    saved = []
    res = parse_docx(buf.getvalue(), save_image=lambda blob, ext: saved.append((blob, ext)) or f"/img/{len(saved)}.{ext}")
    q1, q2 = res["questions"]
    assert res["errors"] == []
    assert q1["question"] == r"Giá trị của biểu thức [[math:\frac{1}{2}]] là?"
    assert q1["options"]["A"] == "[[math:0.5]]" and q1["correct_answer"] == "A"
    assert q2["question"] == "Hình nào là hình tròn?\n[[img:/img/1.png]]"
    assert q2["options"]["A"] == "[[img:/img/2.png]]" and saved[0] == (PNG, "png")


def test_mathtype_ole_reported_not_dropped():
    d = docx.Document()
    d.add_paragraph("Câu 1: Tính")
    run = d.paragraphs[-1].add_run()
    run._r.append(parse_xml('<w:object xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
    for t in ("A. 1", "B. 2", "Đáp án: A"):
        d.add_paragraph(t)
    buf = io.BytesIO()
    d.save(buf)
    res = parse_docx(buf.getvalue())
    assert res["questions"] == [] and "MathType" in res["errors"][0]["reason"]


def test_import_word_saves_images_api():
    T = _h("an.nguyen@studyonline.vn")
    d = docx.Document()
    d.add_paragraph("Câu 1: Ảnh dưới đây là gì?")
    d.add_paragraph().add_run().add_picture(io.BytesIO(PNG))
    for t in ("A. Một điểm", "B. Một đường", "Đáp án: A"):
        d.add_paragraph(t)
    buf = io.BytesIO()
    d.save(buf)
    ctype = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    q = client.post("/api/quizzes/import-word", headers=T,
                    files={"file": ("de.docx", buf.getvalue(), ctype)}).json()["data"]["questions"][0]
    url = q["question"].split("[[img:")[1].split("]]")[0]
    assert url.startswith("/uploads/images/quiz/") and url.endswith(".png")
    got = client.get(url)
    assert got.status_code == 200 and got.content == PNG
    # Dọn file ảnh thử (bản nháp chưa lưu nên không câu hỏi nào dùng)
    from app.core.config import settings

    os.remove(os.path.join(settings.upload_dir, url.removeprefix("/uploads/")))
